"""报告真实文件导出：txt / md / docx / pdf。

统一以 ReportDoc 形状的 dict（Report.payload）为输入，渲染**完整结构化内容**
（摘要 + 财务表 + 章节 + 风险 + 来源 + 免责），不只导出摘要。
- docx：python-docx（封面信息 + 摘要 + 财务表 + 章节 + 风险列表 + 来源 + 免责）。
- pdf：reportlab Platypus（注册 STSong-Light CID 字体保证中文正常渲染，不依赖 weasyprint/GTK）。
- txt / md：同一份结构化报告纯文本/Markdown。
模型 API Key 绝不写入导出文件。
"""

from __future__ import annotations

import datetime as dt
import io
from typing import Any


def _normalize(doc: dict) -> dict:
    fins = doc.get("financials") or {}
    return {
        "ticker": doc.get("ticker", "") or "",
        "name": doc.get("name", "") or "",
        "template": doc.get("template", "equity_deep"),
        "summary": doc.get("summary", "") or "",
        "catalyst": doc.get("catalyst", "") or "",
        "years": list(fins.get("years", []) or []),
        "rows": list(fins.get("rows", []) or []),
        "sections": list(doc.get("sections", []) or []),
        "risks": list(doc.get("risks", []) or []),
        "sources": list(doc.get("sources", []) or []),
        "disclaimer": doc.get("disclaimer", "") or "",
        "as_of": doc.get("as_of", "") or "",
    }


def safe_filename(doc: dict, ext: str) -> str:
    d = _normalize(doc)
    base = (d["name"] or d["ticker"] or "JingYuan-Report").strip() or "report"
    for ch in '\\/:*?"<>|\t\n':
        base = base.replace(ch, "_")
    return f"{base}.{ext}"


# ────────────────────────────────────────────────────────────
# TXT
# ────────────────────────────────────────────────────────────
def render_txt(doc: dict) -> str:
    d = _normalize(doc)
    L: list[str] = [
        f"{d['name'] or d['ticker']} 投资研究报告",
        "=" * 40,
        f"标的：{d['ticker']}    报告日期：{d['as_of'] or dt.date.today().isoformat()}",
        "",
        "【投资摘要】",
        d["summary"],
        "",
    ]
    if d["catalyst"]:
        L += ["【催化剂】", d["catalyst"], ""]
    if d["rows"]:
        L.append("【财务概览】")
        L.append("指标".ljust(12) + "".join(str(y).rjust(14) for y in d["years"]))
        for r in d["rows"]:
            L.append(str(r.get("name", "")).ljust(12) + "".join(str(v).rjust(14) for v in r.get("values", [])))
        L.append("")
    for s in d["sections"]:
        L += [f"【{s.get('title', '')}】", s.get("body", ""), ""]
    if d["risks"]:
        L.append("【风险提示】")
        L += [f"· {r}" for r in d["risks"]]
        L.append("")
    if d["sources"]:
        L.append("【引用来源】")
        for i, s in enumerate(d["sources"], 1):
            L.append(f"[{i}] {s.get('title', '')} · {s.get('date', '')} · {s.get('url', '')}")
        L.append("")
    L += ["-" * 40, d["disclaimer"]]
    return "\n".join(L)


# ────────────────────────────────────────────────────────────
# Markdown（.md 文件下载用，Markdown 在这里是合适的）
# ────────────────────────────────────────────────────────────
def render_md(doc: dict) -> str:
    d = _normalize(doc)
    L: list[str] = [f"# {d['name'] or d['ticker']} 投资研究报告\n"]
    L.append(f"> 标的 **{d['ticker']}** · 报告日期 **{d['as_of']}**\n")
    L.append("## 投资摘要\n" + d["summary"] + "\n")
    if d["catalyst"]:
        L.append("## 催化剂\n" + d["catalyst"] + "\n")
    if d["rows"]:
        L.append("## 财务概览\n")
        L.append("| 指标 | " + " | ".join(str(y) for y in d["years"]) + " |")
        L.append("|" + " --- |" * (len(d["years"]) + 1))
        for r in d["rows"]:
            L.append("| " + str(r.get("name", "")) + " | " + " | ".join(str(v) for v in r.get("values", [])) + " |")
        L.append("")
    for s in d["sections"]:
        L.append(f"## {s.get('title', '')}\n" + s.get("body", "") + "\n")
    if d["risks"]:
        L.append("## 风险提示\n" + "\n".join(f"- {r}" for r in d["risks"]) + "\n")
    if d["sources"]:
        L.append("## 引用来源\n" + "\n".join(
            f"{i}. [{s.get('title', '')}]({s.get('url', '')}) · {s.get('date', '')}"
            for i, s in enumerate(d["sources"], 1)
        ) + "\n")
    L.append("---\n" + d["disclaimer"])
    return "\n".join(L)


# ────────────────────────────────────────────────────────────
# DOCX（python-docx）
# ────────────────────────────────────────────────────────────
def render_docx(doc: dict) -> bytes:
    from docx import Document

    d = _normalize(doc)
    document = Document()
    document.add_heading(f"{d['name'] or d['ticker']} 投资研究报告", level=0)
    document.add_paragraph(f"标的：{d['ticker']}    报告日期：{d['as_of']}")
    document.add_heading("投资摘要", level=1)
    document.add_paragraph(d["summary"])
    if d["catalyst"]:
        document.add_heading("催化剂", level=1)
        document.add_paragraph(d["catalyst"])
    if d["rows"]:
        document.add_heading("财务概览", level=1)
        table = document.add_table(rows=1, cols=len(d["years"]) + 1)
        try:
            table.style = "Light Grid Accent 1"
        except Exception:  # noqa: BLE001
            pass
        hdr = table.rows[0].cells
        hdr[0].text = "指标"
        for i, y in enumerate(d["years"]):
            hdr[i + 1].text = str(y)
        for r in d["rows"]:
            cells = table.add_row().cells
            cells[0].text = str(r.get("name", ""))
            for i, v in enumerate(r.get("values", [])):
                cells[i + 1].text = str(v)
    for s in d["sections"]:
        document.add_heading(str(s.get("title", "")), level=1)
        for para in str(s.get("body", "")).split("\n"):
            if para.strip():
                document.add_paragraph(para)
    if d["risks"]:
        document.add_heading("风险提示", level=1)
        for r in d["risks"]:
            try:
                document.add_paragraph(str(r), style="List Bullet")
            except Exception:  # noqa: BLE001
                document.add_paragraph(f"· {r}")
    if d["sources"]:
        document.add_heading("引用来源", level=1)
        for i, s in enumerate(d["sources"], 1):
            document.add_paragraph(f"[{i}] {s.get('title', '')} · {s.get('date', '')} · {s.get('url', '')}")
    p = document.add_paragraph(d["disclaimer"])
    try:
        p.runs[0].italic = True
    except Exception:  # noqa: BLE001
        pass
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ────────────────────────────────────────────────────────────
# PDF（reportlab + STSong-Light CID 字体支持中文）
# ────────────────────────────────────────────────────────────
_CJK_REGISTERED = False


def _ensure_cjk_font():
    global _CJK_REGISTERED
    if _CJK_REGISTERED:
        return "STSong-Light"
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        _CJK_REGISTERED = True
        return "STSong-Light"
    except Exception:  # noqa: BLE001
        return "Helvetica"  # 退化（中文可能乱码，但不崩）


def render_pdf(doc: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    d = _normalize(doc)
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm, leftMargin=16 * mm, rightMargin=16 * mm
    )
    cjk = _ensure_cjk_font()
    base = getSampleStyleSheet()
    st_title = ParagraphStyle("ZhTitle", parent=base["Title"], fontName=cjk, fontSize=18, leading=24)
    st_h1 = ParagraphStyle("ZhH1", parent=base["Heading1"], fontName=cjk, fontSize=14, leading=20, spaceBefore=8, spaceAfter=4)
    st_body = ParagraphStyle("ZhBody", parent=base["BodyText"], fontName=cjk, fontSize=10, leading=16, wordWrap="CJK")
    st_disc = ParagraphStyle("ZhDisc", parent=st_body, fontSize=8, textColor=colors.grey)

    def esc(t: Any) -> str:
        return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")

    story: list = []
    story.append(Paragraph(esc(f"{d['name'] or d['ticker']} 投资研究报告"), st_title))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(esc(f"标的：{d['ticker']}    报告日期：{d['as_of']}"), st_body))
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph("投资摘要", st_h1))
    story.append(Paragraph(esc(d["summary"]), st_body))
    if d["catalyst"]:
        story.append(Paragraph("催化剂", st_h1))
        story.append(Paragraph(esc(d["catalyst"]), st_body))
    if d["rows"]:
        story.append(Paragraph("财务概览", st_h1))
        data = [["指标"] + [str(y) for y in d["years"]]]
        for r in d["rows"]:
            data.append([str(r.get("name", ""))] + [str(v) for v in r.get("values", [])])
        tbl = Table(data, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), cjk),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9c46a")),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 3 * mm))
    for s in d["sections"]:
        story.append(Paragraph(esc(s.get("title", "")), st_h1))
        story.append(Paragraph(esc(s.get("body", "")), st_body))
    if d["risks"]:
        story.append(Paragraph("风险提示", st_h1))
        for r in d["risks"]:
            story.append(Paragraph("• " + esc(r), st_body))
    if d["sources"]:
        story.append(Paragraph("引用来源", st_h1))
        for i, s in enumerate(d["sources"], 1):
            story.append(Paragraph(f"[{i}] " + esc(f"{s.get('title', '')} · {s.get('date', '')} · {s.get('url', '')}"), st_body))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(esc(d["disclaimer"]), st_disc))
    pdf.build(story)
    return buf.getvalue()
